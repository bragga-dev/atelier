"""
Order Export Service — gera a planilha de pedidos (CSV ou XLSX) pro admin
separar/embalar. Cada linha é um pedido; os itens vêm resumidos numa
coluna só ("2x Pulseira de Cristais; 1x Colar de Búzios").
"""
import csv
import io
from typing import Iterable, List

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from atelier.apps.payments.models.order_model import Order

EXPORT_HEADERS = [
    "Código",
    "Data",
    "Status",
    "Cliente",
    "E-mail",
    "Itens",
    "Subtotal",
    "Frete",
    "Total",
    "CEP",
    "Endereço",
    "Bairro",
    "Cidade",
    "Estado",
]


def _customer_name(order: Order) -> str:
    client = getattr(order.user_id, "client_profile", None)
    if not client:
        return ""
    return f"{client.first_name or ''} {client.last_name or ''}".strip()


def _items_summary(order: Order) -> str:
    return "; ".join(f"{item.order_item_quantity}x {item.product_id.product_name}" for item in order.items.all())


def _order_row(order: Order) -> list:
    address = order.shipping_address
    street = f"{address.street}, {address.number}"
    if address.complement:
        street += f" - {address.complement}"

    return [
        order.code,
        order.created_at.strftime("%d/%m/%Y %H:%M"),
        order.get_order_status_display(),
        _customer_name(order),
        order.user_id.email,
        _items_summary(order),
        f"{order.subtotal:.2f}",
        f"{order.order_shipping_total:.2f}",
        f"{order.total_geral:.2f}",
        address.cep,
        street,
        address.neighborhood,
        address.city,
        address.state,
    ]


def build_orders_rows(orders: Iterable[Order]) -> List[list]:
    return [_order_row(order) for order in orders]


def render_orders_csv(orders: Iterable[Order]) -> bytes:
    buffer = io.StringIO()
    # BOM (\ufeff) — sem isso o Excel no Windows abre acentuação quebrada.
    buffer.write("\ufeff")
    writer = csv.writer(buffer, delimiter=";")
    writer.writerow(EXPORT_HEADERS)
    writer.writerows(build_orders_rows(orders))
    return buffer.getvalue().encode("utf-8")


def render_orders_xlsx(orders: Iterable[Order]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Pedidos"

    rows = build_orders_rows(orders)

    ws.append(EXPORT_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append(row)

    for i, header in enumerate(EXPORT_HEADERS, start=1):
        column_values = [header] + [str(row[i - 1]) for row in rows]
        width = min(max(len(v) for v in column_values) + 2, 50)
        ws.column_dimensions[get_column_letter(i)].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()